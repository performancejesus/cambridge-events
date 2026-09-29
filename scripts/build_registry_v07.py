"""Этап 6c: data/cambridge_event_sources_v0.7.xlsx из v0.6 (v0.6 не изменяется).

Новые и пересмотренные источники (data/p6c_decisions.json) — решение в «Заметках», цвет ID; ошибочные решения этапов
5–6 отмечены. Листы: «Этап 6c» (источник → новых событий → решение; таблица по спорту — в docs/stage6c_report.md),
«Не разобрано» (таблица unparsed_sources: ID, URL, тип проблемы, с какой даты, что теряем, что делать, проверено).
"""

from __future__ import annotations

import json
import shutil
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from build_registry_v04 import GREEN, HEADER_FILL, HEADER_FONT, RED, YELLOW  # noqa: E402
from pipeline import unparsed  # noqa: E402
from pipeline.db import connect  # noqa: E402

SRC = ROOT / "data" / "cambridge_event_sources_v0.6.xlsx"
DST = ROOT / "data" / "cambridge_event_sources_v0.7.xlsx"
FILL = {"подключён": GREEN, "закрыт": RED, "не сезон": YELLOW, "не нужен": YELLOW}
NEW = {"S149": ("King's College Choir — concerts", "Площадка", "https://kingscollegechoir.com/concert/", "Концерты", "Центр", 2),
       "S150": ("Cambridge Music (cmp.cam.ac.uk) — West Road, CUMS, колледжи", "Агрегатор", "https://www.cmp.cam.ac.uk/events/", "Концерты", "Центр", 2),
       "S151": ("Music Live Cambridge", "Агрегатор", "https://musiclivecambridge.com/", "Концерты", "Центр", 3),
       "S152": ("Find a Race — забеги в 40 км", "Агрегатор / регистрация", "https://findarace.com/events", "Спорт — участвовать", "Час", 3),
       "S153": ("Royston Museum", "Музей", "https://www.roystonmuseum.org.uk/what-s-on", "Выставки и музеи", "До 30 мин", 3),
       "S154": ("Cambridge United Women", "Клуб", "https://www.cambridgeunited.com/fixture/list/844", "Спорт — смотреть", "Центр", 2),
       "S155": ("Peterborough Phantoms (хоккей, NIHL)", "Клуб", "https://www.gophantoms.co.uk/gameday/2026-27-fixtures", "Спорт — смотреть", "Кембриджшир", 2),
       "S156": ("Vue Cambridge (The Grafton)", "Кино", "https://www.myvue.com/cinema/cambridge/whats-on", "Кино", "Центр", 2),
       "S157": ("Light Cinema Cambridge", "Кино", "https://www.thelight.co.uk/cinema/cambridge", "Кино", "Центр", 2),
       "S158": ("Saffron Screen", "Кино", "https://www.saffronscreen.com/", "Кино", "До 30 мин", 3),
       "S159": ("Babylon Cinema, Ely", "Кино", "https://www.babylonely.co.uk/", "Кино", "До 30 мин", 3),
       "S160": ("CURUFC — университетское регби", "Клуб", "https://curufc.com/events-tickets/fixtures/", "Спорт — смотреть", "Центр", 3),
       "S161": ("FA Full-Time", "Лига", "https://fulltime.thefa.com/", "Спорт — смотреть", "Графство", 3),
       "S163": ("Heong Gallery (Downing College)", "Галерея", "https://www.dow.cam.ac.uk/creative-arts/heong-gallery", "Выставки и музеи", "Центр", 2),
       "S164": ("The Women's Art Collection (Murray Edwards)", "Галерея", "https://www.murrayedwards.cam.ac.uk/womens-art-collection", "Выставки и музеи", "Центр", 2),
       "S165": ("Разбор подборок («Things to do…»)", "Производный", "", "Сквозные агрегаторы", "Центр", 3),
       "S166": ("Trinity College events", "Колледж", "https://www.trin.cam.ac.uk/events/", "Лекции", "Центр", 3)}


def main() -> None:
    today = date.today().isoformat()
    dec = {k: v for k, v in json.loads((ROOT / "data" / "p6c_decisions.json").read_text()).items() if not k.startswith("_")}
    uniq_path = ROOT / "docs" / "stage6c_unique.json"
    uniq = json.loads(uniq_path.read_text()) if uniq_path.exists() else {}
    shutil.copy(SRC, DST)
    wb = openpyxl.load_workbook(DST)
    ws = wb["Источники"]
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr)}
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    report = []
    for sid, (state, decision, note, kind, wrong) in dec.items():
        if sid not in rows:
            nm = NEW.get(sid, (sid, "", "", "", "", 3))
            ws.append([sid, nm[3], nm[0], nm[1], nm[2], "", nm[4], "", "Ежедневно", nm[5]])
            rows[sid] = ws.max_row
        r = rows[sid]
        old = ws.cell(r, col["Заметки"]).value or ""
        text = f"Этап 6c: {decision} — {note}" + (f" Ошибочное прежнее решение: {wrong}." if wrong else "")
        ws.cell(r, col["Заметки"], (old + " " if old else "") + text)
        ws.cell(r, col["Проверено"], today)
        ws.cell(r, col["ID"]).fill = FILL.get(decision, PatternFill(fill_type=None))
        u = uniq.get(sid, {})
        report.append([sid, ws.cell(r, col["Источник"]).value, state, decision, u.get("events", "—"), u.get("new", "—"),
                       u.get("window", "—"), note, wrong])
    for name in ("Этап 6c", "Не разобрано"):
        if name in wb.sheetnames:
            del wb[name]
    rep = wb.create_sheet("Этап 6c")
    rep.append(["ID", "Источник", "Состояние", "Решение", "Событий от источника", "Новых уникальных", "Новых в окне 1–11.10",
                "Примечание", "Какое прежнее решение было ошибочным и почему"])
    for c in rep[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for row in sorted(report, key=lambda x: (x[3] != "подключён", x[0])):
        rep.append(row)
        rep.cell(rep.max_row, 4).fill = FILL.get(row[3], PatternFill(fill_type=None))
    for letter, width in zip("ABCDEFGHI", (8, 45, 12, 12, 12, 12, 12, 80, 80)):
        rep.column_dimensions[letter].width = width
    con = connect()
    unparsed.init(con)
    un = wb.create_sheet("Не разобрано")
    un.append(["ID", "Источник", "URL", "Тип проблемы", "Подробности", "С какой даты", "Что теряем", "Что делать",
               "Последняя проверка", "Статус"])
    for c in un[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for r in con.execute("SELECT * FROM unparsed_sources ORDER BY status, key"):
        un.append([r["key"], r["name"], r["url"], unparsed.PROBLEM_RU.get(r["problem"], r["problem"]), r["detail"],
                   r["since"], r["losing"], r["action"], (r["last_checked"] or "")[:10],
                   "открыто" if r["status"] == "open" else "решено"])
    for letter, width in zip("ABCDEFGHIJ", (22, 40, 50, 28, 50, 12, 45, 35, 12, 10)):
        un.column_dimensions[letter].width = width
    help_ws = wb["Как пользоваться"]
    help_ws.append([f"v0.7 ({today}): этап 6c — коллекторы из решений после 6b, зрительский спорт, кино, колледжи; "
                    "правило robots.txt 4xx (RFC 9309); листы «Этап 6c» и «Не разобрано» (перепроверка раз в неделю)."])
    help_ws.cell(help_ws.max_row, 1).font = Font(italic=True)
    wb.save(DST)
    print(DST.relative_to(ROOT), len(report), con.execute("SELECT count(*) FROM unparsed_sources").fetchone()[0])


if __name__ == "__main__":
    main()
