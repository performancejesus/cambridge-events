"""Этап 7d: data/cambridge_event_sources_v0.11.xlsx из v0.10 (v0.10 не изменяется).

- Лист «Источники»: места этапа 7d — National Trust (S070) и Audley End (S139) подключены по открытым данным страниц;
  IWM Duxford (S053) и Shepreth (S135) — закрыты, альтернативы; новые S178–S183 (Bury Lane Farm Shop, Museum of
  Technology, Ely Museum, Corn Exchange King's Lynn, кинотеатры зоны с городом, Visit West Norfolk); фестивали (S031,
  S033, S147) — статус в recurring_events; musiclivecambridge (S151) — склейка дублей.
- Лист «Не разобрано» — свежая выгрузка unparsed_sources.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from build_registry_v04 import HEADER_FILL, HEADER_FONT  # noqa: E402
from pipeline import unparsed  # noqa: E402
from pipeline.db import connect  # noqa: E402

SRC = ROOT / "data" / "cambridge_event_sources_v0.10.xlsx"
DST = ROOT / "data" / "cambridge_event_sources_v0.11.xlsx"
T = "2026-09-30"
NOTES = {
    "S070": "Этап 7d: подключён — страницы событий Wimpole, Anglesey Abbey, Wicken Fen отвечают 200 честному "
            "User-Agent (robots.txt разрешает); события — из данных страницы (__NEXT_DATA__: eventsPageData), без "
            "модели. Если снова включится Radware — не обходим, возвращаемся к ручному календарю.",
    "S139": "Этап 7d: подключён — открытый API поиска событий English Heritage (/api/eventsearch, место 34166 — Audley "
            "End), тот же, что вызывает скрипт страницы; события для членов EH — с пометкой доступа.",
    "S053": "Этап 7d: по-прежнему 403 с проверкой бот-защиты (iwm.org.uk) — не обходим. Альтернативный путь: события "
            "Duxford, которые публикует Visit Cambridge (S001), и статьи газет; крупные авиашоу — вручную в recurring_events.",
    "S135": "Этап 7d: 403 / обрыв соединения, иногда пропускает (нестабильно) — не обходим, лист «Не разобрано».",
    "S031": "Этап 7d: статус в recurring_events (R09): на сайте только дата окончания (1 августа 2027) — начало вручную.",
    "S033": "Этап 7d: R07 — дата 2027 ещё не объявлена (ожидаем, июнь).",
    "S147": "Этап 7d: R22 — дата берётся только рядом с «Stourbridge Fair» (на странице CPPF календарь всех событий "
            "капеллы; в 7c ошибочно взялась дата «Гамлета»). R24 — только огни Кембриджа (в 7c совпало с Висбечем).",
    "S151": "Этап 7d: дубли склеиваются с событием площадки по slug страницы + дате + площадке (разрешено владельцем; "
            "снимок базы до склейки): 7 пар при первом запуске.",
}
NEW = [  # ID, категория, источник, тип, URL, endpoint, покрывает, способ, частота, robots, заметка, зона
    ("S178", "Семьи", "Bury Lane Farm Shop (Melbourn)", "Ферма (HTML)", "https://burylane.co.uk/events/",
     "https://burylane.co.uk/events/", "сезонные события фермы: pick-your-own, ярмарки, Рождество, семейные дни",
     "Видимый текст → Haiku (кэш по хэшу страницы)", "Еженедельно", "разрешено",
     "Этап 7d: адрес сайта — burylane.co.uk (burylanefarmshop.co.uk — старый).", "До 30 мин"),
    ("S179", "Музеи", "Cambridge Museum of Technology", "Музей (HTML)", "https://www.museumoftechnology.com/whats-on",
     "https://www.museumoftechnology.com/whats-on", "дни паровых машин, семейные занятия, экскурсии, лекции",
     "Видимый текст → Haiku", "Еженедельно", "разрешено", "Этап 7d.", "Центр"),
    ("S180", "Музеи", "Ely Museum", "Музей (HTML)", "https://www.elymuseum.org.uk/whats-on-elymuseum/",
     "https://www.elymuseum.org.uk/whats-on-elymuseum/", "выставки, лекции, семейные занятия",
     "Видимый текст → Haiku", "Еженедельно", "разрешено", "Этап 7d.", "До 30 мин"),
    ("S181", "Театр", "Corn Exchange King's Lynn (Alive West Norfolk)", "Площадка (HTML)",
     "https://www.kingslynncornexchange.co.uk/", "https://www.kingslynncornexchange.co.uk/",
     "театр, комедия, концерты (ближайшие показы на главной)", "Видимый текст → Haiku", "Еженедельно", "разрешено",
     "Этап 7d: список /theatre/whats-on/ подгружается POST-запросом скрипта — берём главную; кино — в S182. Кингс-Линн — "
     "исключение географии (как West Suffolk): зона «до часа» в радиусе 6 км от центра.", "До часа"),
    ("S182", "Кино", "Кинотеатры зоны (с пометкой city)", "Кинотеатры (HTML / JSON)",
     "https://wisbech.thelight.co.uk", "мини-гид Light (JSON); остальные — видимый текст → Haiku",
     "расписание / «сейчас в прокате»: Light Wisbech, Royal (Merlin) и Screen St Ives, Majestic и Corn Exchange "
     "King's Lynn, Peterborough Arts Cinema, Luxe Wisbech, Haverhill Arts Centre",
     "regional_showings (все фильмы с городом); события — только трансляции и спецпоказы", "Еженедельно",
     "разрешено (закрытые — лист «Не разобрано»)",
     "Этап 7d: «собирать шире, публиковать уже» — в кембриджский выпуск только показы, которых нет в Light / Arts "
     "Picturehouse (проверка 38). Сети (Cineworld, Odeon), Abbeygate, Saffron Screen — 403 / бот-защита; Babylon (Ely "
     "Maltings), Key Theatre, Kings Newmarket — обрыв соединения через шлюз.", "До часа"),
    ("S183", "Город", "Visit West Norfolk — события Кингс-Линна и округи", "Туристический сайт (HTML)",
     "https://www.visitwestnorfolk.com/experiences/events/", "https://www.visitwestnorfolk.com/experiences/events/",
     "фестивали, ярмарки, лекции, выставки (в т. ч. Stories of Lynn / Lynn Museum)", "Видимый текст → Haiku",
     "Еженедельно", "разрешено",
     "Этап 7d: сайт то пускает, то отдаёт заглушку Sucuri — в этом случае прогон источника — ошибка (не «событий нет»), "
     "не обходим. Lynn Museum (Norfolk Museums — 403) и True's Yard (Sucuri) — через этот календарь.", "До часа"),
]



def main() -> None:
    shutil.copy(SRC, DST)
    wb = openpyxl.load_workbook(DST)
    ws = wb["Источники"]
    hdr = [c.value for c in ws[1]]
    ci = {h: hdr.index(h) + 1 for h in hdr if h}
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
    for sid, note in NOTES.items():
        r = rows[sid]
        ws.cell(r, ci["Заметки"], ((ws.cell(r, ci["Заметки"]).value or "") + " " + note).strip())
        ws.cell(r, ci["Проверено"], T)
    ws.cell(rows["S070"], ci["Способ сбора"], "HTML (данные страницы __NEXT_DATA__)")
    ws.cell(rows["S139"], ci["Способ сбора"], "JSON API (eventsearch)")
    for sid, cat, name, typ, url, endpoint, covers, how, freq, robots, note, zone in NEW:
        if sid in rows:
            continue
        ws.append([None] * len(hdr))
        r = ws.max_row
        for k, v in {"ID": sid, "Категория": cat, "Источник": name, "Тип источника": typ, "URL": url,
                     "Endpoint для сбора": endpoint, "Что покрывает": covers, "Зона": zone, "Способ сбора": how,
                     "Частота": freq, "Приоритет": 2, "URL проверен": "да", "Заметки": note, "robots.txt": robots,
                     "Проверено": T}.items():
            ws.cell(r, ci[k], v)
    if "Не разобрано" in wb.sheetnames:
        del wb["Не разобрано"]
    un = wb.create_sheet("Не разобрано")
    un.append(["ID", "Источник", "URL", "Тип проблемы", "Подробности", "С какой даты", "Что теряем", "Что делать",
               "Последняя проверка", "Статус"])
    for c in un[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    con = connect()
    unparsed.init(con)
    for r in con.execute("SELECT * FROM unparsed_sources ORDER BY status, key"):
        un.append([r["key"], r["name"], r["url"], unparsed.PROBLEM_RU.get(r["problem"], r["problem"]), r["detail"],
                   r["since"], r["losing"], r["action"], (r["last_checked"] or "")[:10],
                   "открыто" if r["status"] == "open" else "решено"])
    for letter, width in zip("ABCDEFGHIJ", (22, 40, 50, 28, 50, 12, 45, 35, 12, 10)):
        un.column_dimensions[letter].width = width
    help_ws = wb["Как пользоваться"]
    help_ws.append([f"v0.11 ({T}): этап 7d — музеи, усадьбы, фермы (S070, S139, S178–S180), Кингс-Линн (S181, S183), "
                    "кинотеатры зоны с городом (S182); закрытые места — лист «Не разобрано»."])
    help_ws.cell(help_ws.max_row, 1).font = Font(italic=True)
    wb.save(DST)
    print(DST.relative_to(ROOT))


if __name__ == "__main__":
    main()
