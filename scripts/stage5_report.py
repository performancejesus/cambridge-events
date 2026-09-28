"""Этап 5: отчёт — источники P2 и расширение географии → docs/stage5_report.md.

Новые уникальные события источника — будущие события в зоне, у которых нет ни одного источника P1 (то, чего не было
до этапа 5). Район — admin_district по postcode (postcodes.io) или по населённому пункту площадки.
"""

from __future__ import annotations

import json
import sqlite3
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from pipeline.db import connect  # noqa: E402

LISTED = ("центр", "до 30 мин", "до часа", "Кембриджшир, дальше часа")
CONTROLS = [("parkrun", "parkrun"), ("Cambridgeshire Cats", "Cambridgeshire Cats"), ("IWM Duxford", "Duxford"),
            ("Kettle's Yard", "Kettle"), ("Saffron Hall", "Saffron Hall")]


def p1_ids() -> set[str]:
    ws = openpyxl.load_workbook(ROOT / "data" / "cambridge_event_sources_v0.4.xlsx")["Источники"]
    return {r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[9] == 1}


def district(con: sqlite3.Connection, e: sqlite3.Row) -> str:
    if e["postcode"]:
        r = con.execute("SELECT admin_district FROM postcodes WHERE replace(postcode,' ','')=replace(?,' ','')",
                        (e["postcode"],)).fetchone()
        if r and r[0]:
            return r[0]
    if e["venue_id"]:
        v = con.execute("SELECT lat, lon, precision FROM venues WHERE venue_id=?", (e["venue_id"],)).fetchone()
        if v and v["precision"] == "place":
            p = con.execute("SELECT district, county FROM places WHERE lat=? AND lon=?", (v["lat"], v["lon"])).fetchone()
            if p:
                return p["district"] or p["county"]
    return "Cambridge (адрес не уточнён)" if e["address_unknown"] else "не определён"


def main() -> None:
    con = connect()
    today = date.today().isoformat()
    p1 = p1_ids() - {"S051"}          # Kettle's Yard (P1) получил коллектор только на этапе 5 — считаем как новый
    dec = {k: v for k, v in json.loads((ROOT / "data" / "p2_decisions.json").read_text()).items() if not k.startswith("_")}
    names = json.loads((ROOT / "data" / "p2_candidates.json").read_text())["_names"]
    ws = openpyxl.load_workbook(ROOT / "data" / "cambridge_event_sources_v0.4.xlsx")["Источники"]
    names |= {r[0]: [r[2], r[3], r[6]] for r in ws.iter_rows(min_row=2, values_only=True) if r[0] and r[0] not in names}

    events = con.execute("""SELECT e.*, (SELECT group_concat(DISTINCT source_id) FROM event_sources s
        WHERE s.event_id = e.event_id) AS srcs FROM events e
        WHERE coalesce(e.date_end, e.date_start) >= ? AND e.status NOT IN ('past', 'cancelled')""", (today,)).fetchall()
    unique, found, out_zone = Counter(), Counter(), Counter()
    by_district, by_district_new = Counter(), Counter()
    for e in events:
        srcs = set((e["srcs"] or "").split(",")) - {""}
        new_srcs = {s for s in srcs if s in dec}
        if e["zone"] == "out_of_zone":
            for s in new_srcs:
                out_zone[s] += 1
            continue
        if e["zone"] not in LISTED:
            continue
        dist = district(con, e)
        by_district[dist] += 1
        for s in new_srcs:
            found[s] += 1
        if new_srcs and not srcs & p1:
            by_district_new[dist] += 1
            for s in new_srcs:
                unique[s] += 1

    L = [f"# Этап 5 — источники приоритета 2 и расширение географии ({today})", "",
         "Проверка: `scripts/probe_sources.py --set p2` → `data/probe_results_p2.json` (кандидаты — `data/p2_candidates.json`).",
         "Решения по источникам — `data/p2_decisions.json`. Новые уникальные события — будущие события в зоне, у которых "
         "нет ни одного источника P1.", "", "## Итог проверки", ""]
    states, decisions = Counter(v[0] for v in dec.values()), Counter(v[1] for v in dec.values())
    L += [f"- Источников проверено: {len(dec)} (32 из реестра с приоритетом 2, 30 новых S098–S127, 2 страницы "
          f"агрегаторов S128/S129, Kettle's Yard S051 из P1).",
          f"- Состояние: " + ", ".join(f"{k} — {v}" for k, v in states.most_common()) + ".",
          f"- Решение: " + ", ".join(f"{k} — {v}" for k, v in decisions.most_common()) + ".", "",
          "| ID | Источник | Состояние | Решение | Новых уникальных | Всего в зоне | Вне зоны | Примечание |",
          "|---|---|---|---|---|---|---|---|"]
    order = sorted(dec, key=lambda s: (-unique[s], {"подключён": 0}.get(dec[s][1], 1), s))
    for sid in order:
        st, d, note = dec[sid]
        nm = names.get(sid, [sid])[0]
        L.append(f"| {sid} | {nm} | {st} | {d} | {unique[sid] or '—'} | {found[sid] or '—'} | {out_zone[sid] or '—'} | {note} |")

    L += ["", "## События по районам (будущие, в зоне)", "",
          "| Район | Всего событий | из них новых (этап 5) |", "|---|---|---|"]
    for dist, n in by_district.most_common():
        L.append(f"| {dist} | {n} | {by_district_new[dist] or '—'} |")

    L += ["", "## Контрольные позиции", "", "| Что | В базе (будущие) | Откуда / почему нет |", "|---|---|---|"]
    why = {"parkrun": "parkrun.org.uk: robots.txt отвечает 403 — считаем запретом; регулярные забеги — только вручную",
           "Cambridgeshire Cats": "сайт живой (WordPress, новости), расписания на сайте нет; сезон BAFA — весна–лето",
           "IWM Duxford": "iwm.org.uk — 403 (бот-защита) на все страницы; в Ents24/Skiddle не найден",
           "Kettle's Yard": "коллектор S051 (этап 5): страницы событий",
           "Saffron Hall": "коллектор S017 (этап 5) + Ents24 по Saffron Walden (S128)"}
    for label, pat in CONTROLS:
        n = con.execute("""SELECT count(*) FROM events WHERE (title LIKE ? OR venue_name LIKE ?)
            AND coalesce(date_end, date_start) >= ?""", (f"%{pat}%", f"%{pat}%", today)).fetchone()[0]
        L.append(f"| {label} | {n or 'нет'} | {why[label]} |")

    L += ["", "## Вне зоны", "",
          f"- Помечено `out_of_zone` (будущие): " + ", ".join(f"{s} — {n}" for s, n in out_zone.most_common()) +
          ". Страница города у Skiddle включает соседние регионы — зона считается по postcode площадки, в выпуск не идут.", ""]
    # решения после этапа 4b: оценка важности (топ-15 периода первого выпуска и три проверочных события)
    L += ["## Решения после этапа 4b: оценка важности", "",
          "Wikipedia и известность — только за того, кто на сцене (поле `performer`); трибьюты, шоу «по мотивам» и "
          "составы с одним известным именем — с потолком; у постановок классики считается труппа, не пьеса; бесплатно / "
          "цена неизвестна и площадка без вместимости — нейтральный сигнал; лекции без билетов — известность × 1.6; "
          "пересчёт — только при изменении входных данных (сигнатура); «Главное на выходные» — 3–5 лучших, не ниже 4.", "",
          "| # | Оценка | Дата | Событие | Главные причины |", "|---|---|---|---|---|"]
    seen: set[str] = set()
    for e in con.execute("""SELECT title, date_start, importance_score, importance_reason FROM events
            WHERE date_start <= '2026-10-11' AND coalesce(date_end, date_start) >= '2026-09-28'
            AND status NOT IN ('past', 'cancelled') AND zone != 'out_of_zone' AND importance_score IS NOT NULL
            ORDER BY importance_score DESC, date_start"""):
        if e["title"] in seen:
            continue
        seen.add(e["title"])
        reason = "; ".join(e["importance_reason"].split("; ")[:3])
        L.append(f"| {len(seen)} | {e['importance_score']:g} | {e['date_start']} | {e['title'][:70]} | {reason} |")
        if len(seen) == 15:
            break
    L += ["", "Проверочные события:", ""]
    for pat in ("Quantum gravity%", "Boyband In The Buff", "CAST 2026: King Lear"):
        e = con.execute("SELECT title, importance_score, importance_reason FROM events WHERE title LIKE ? ORDER BY date_start LIMIT 1",
                        (pat,)).fetchone()
        L.append(f"- **{e['title'][:60]}** — {e['importance_score']:g}: {e['importance_reason']}")
    notes = ROOT / "docs" / "stage5_notes.md"
    if notes.exists():
        L += ["", notes.read_text().strip()]
    (ROOT / "docs" / "stage5_report.md").write_text("\n".join(L) + "\n")
    print("docs/stage5_report.md")


if __name__ == "__main__":
    main()
