"""Этап 7c: data/cambridge_event_sources_v0.10.xlsx из v0.9 (v0.9 не изменяется).

- Лист «Источники»: Cambridge Union (S049) — где на самом деле публикуется программа; новые S176 (Varsity — статьи о
  Cambridge Union, RSS) и S177 (termcard Cambridge Union на Issuu — страницы-изображения, разбор моделью); S123 — новая
  страница календаря (старая /fixtures отвечает 404); S144 — ссылка на страницу события и цена «Free» со страницы.
- Лист «Не разобрано» — свежая выгрузка unparsed_sources.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from build_registry_v04 import HEADER_FILL, HEADER_FONT  # noqa: E402
from pipeline import unparsed  # noqa: E402
from pipeline.db import connect  # noqa: E402

SRC = ROOT / "data" / "cambridge_event_sources_v0.9.xlsx"
DST = ROOT / "data" / "cambridge_event_sources_v0.10.xlsx"
T = "2026-09-30"
NOTES = {
    "S049": "Этап 7c: на cus.org/event клуб публикует только билетные мероприятия (ужины, вечеринки в The Orator); "
            "программа гостей и дебатов — termcard (S177) и статьи Varsity (S176). В выпуске — одна строка «В Cambridge "
            "Union на этой неделе (для членов клуба): …»; открытые для всех — обычными пунктами после появления на cus.org.",
    "S123": "Этап 7c: ссылка календаря — theposh.com/matches/fixtures (прежняя /fixtures отвечает 404, нашла проверка "
            "ссылок выпуска).",
    "S144": "Этап 7c: ссылка — на страницу события (не на общий список); цена «Free», если страница события или её "
            "«Book now» говорит «free and open to all» (лекция Josephine Crawley Quinn).",
}
NEW = [  # ID, категория, источник, тип, URL, endpoint, покрывает, способ, частота, robots, заметка
    ("S176", "Лекции", "Varsity — статьи о Cambridge Union", "СМИ (RSS)", "https://www.varsity.co.uk/news",
     "http://feeds.varsity.co.uk/varsity/news", "анонсы termcard и гостей Union («… to speak at the Union»)",
     "RSS + Haiku (только статьи о Union)", "Ежедневно", "разрешено; ИИ-агенты не закрыты; robots.txt feeds.varsity — 404",
     "Этап 7c: статья о termcard выходит в первые дни триместра (Michaelmas 2025 — 5 октября), но даты в ней есть не у "
     "всех гостей. Гости и даты — кандидаты; подтверждение — termcard (S177) или cus.org (S049)."),
    ("S177", "Лекции", "Cambridge Union — termcard на Issuu", "Журнал (изображения)",
     "https://issuu.com/thecambridgeunion", "https://image.isu.pub/<id документа>/jpg/page_N.jpg",
     "вся программа триместра: дебаты, гости, панели, светские события, доступ",
     "Изображения страниц → один запрос к модели с изображениями (Haiku)", "Раз в триместр",
     "разрешено (issuu.com: /…/docs/ открыт для всех агентов; image.isu.pub: robots.txt 403 — правил нет)",
     "Этап 7c: проверено на Michaelmas 2025 (84 стр., 30 событий, $0.14), Lent 2026 (88 стр., 41, $0.15), Easter 2026 "
     "(66 стр., 24, $0.11). Termcard выкладывается на Issuu через 1–2 недели после начала триместра (Michaelmas 2025 — "
     "20 октября, Easter 2026 — 29 апреля). YouTube не используем."),
]


def main() -> None:
    shutil.copy(SRC, DST)
    wb = openpyxl.load_workbook(DST)
    ws = wb["Источники"]
    hdr = [c.value for c in ws[1]]
    ci = {h: hdr.index(h) + 1 for h in hdr if h}
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
    for sid, note in NOTES.items():
        r = rows[sid]
        ws.cell(r, ci["Заметки"], ((ws.cell(r, ci["Заметки"]).value or "") + " " + note).strip())
        ws.cell(r, ci["Проверено"], T)
    ws.cell(rows["S123"], ci["URL"], "https://www.theposh.com/matches/fixtures")
    for sid, cat, name, typ, url, endpoint, covers, how, freq, robots, note in NEW:
        if sid in rows:
            continue
        ws.append([None] * len(hdr))
        r = ws.max_row
        for k, v in {"ID": sid, "Категория": cat, "Источник": name, "Тип источника": typ, "URL": url,
                     "Endpoint для сбора": endpoint, "Что покрывает": covers, "Зона": "Центр", "Способ сбора": how,
                     "Частота": freq, "Приоритет": 2, "URL проверен": "да", "Заметки": note, "robots.txt": robots,
                     "Проверено": T}.items():
            ws.cell(r, ci[k], v)
    if "Не разобрано" in wb.sheetnames:
        del wb["Не разобрано"]
    un = wb.create_sheet("Не разобрано")
    un.append(["ID", "Источник", "URL", "Тип проблемы", "Подробности", "С какой даты", "Что теряем", "Что делать",
               "Последняя проверка", "Статус"])
    for c in un[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    con = connect()
    unparsed.init(con)
    for r in con.execute("SELECT * FROM unparsed_sources ORDER BY status, key"):
        un.append([r["key"], r["name"], r["url"], unparsed.PROBLEM_RU.get(r["problem"], r["problem"]), r["detail"],
                   r["since"], r["losing"], r["action"], (r["last_checked"] or "")[:10],
                   "открыто" if r["status"] == "open" else "решено"])
    for letter, width in zip("ABCDEFGHIJ", (22, 40, 50, 28, 50, 12, 45, 35, 12, 10)):
        un.column_dimensions[letter].width = width
    help_ws = wb["Как пользоваться"]
    help_ws.append([f"v0.10 ({T}): этап 7c — Cambridge Union: Varsity (S176) и termcard на Issuu (S177); ссылки S123 и "
                    "S144 исправлены по проверкам выпуска."])
    help_ws.cell(help_ws.max_row, 1).font = Font(italic=True)
    wb.save(DST)
    print(DST.relative_to(ROOT))


if __name__ == "__main__":
    main()
