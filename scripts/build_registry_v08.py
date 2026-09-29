"""Этап 6d: data/cambridge_event_sources_v0.8.xlsx из v0.7 (v0.7 не изменяется).

- Лист «Источники»: колонка «Партнёрская программа» — у источника-продавца (Skiddle, Eventbrite, Jockey Club…) —
  его программа; у остальных — продавцы, к которым ведут события источника в окне выпуска, и есть ли у них программа.
- Лист «Монетизация»: продавец → программа → сеть → комиссия → cookie → email разрешён → пометка → источник условий →
  дата проверки (data/affiliates_6d.json), плюс принципы подключения и итоги (data/monetization_6d.json).
- Лист «Не разобрано» — обновлён из unparsed_sources.
Ни в какие программы не регистрировались; ссылки в выпуске не меняются.
"""

from __future__ import annotations

import json
import shutil
import sys
from collections import Counter
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from build_registry_v04 import HEADER_FILL, HEADER_FONT  # noqa: E402
from monetization import has_programme, programme_of  # noqa: E402
from pipeline import domains, unparsed  # noqa: E402
from pipeline.db import connect  # noqa: E402

SRC = ROOT / "data" / "cambridge_event_sources_v0.7.xlsx"
DST = ROOT / "data" / "cambridge_event_sources_v0.8.xlsx"
PRINCIPLES = [
    "Комиссия не влияет на отбор и порядок пунктов выпуска.",
    "Ссылку не менять на другого продавца ради комиссии: партнёрской может быть только ссылка на того же продавца, "
    "у которого площадка и так продаёт билеты.",
    "Партнёрские ссылки помечать в письме по правилам ASA/CAP: метка «Ad» (в русской версии — «Реклама») у каждой "
    "партнёрской ссылки, до ссылки и заметно; «affiliate» отдельно и «may earn a commission» — недостаточно "
    "(https://www.asa.org.uk/advice-online/affiliate-marketing.html). Отдельных указаний по письмам нет — формулировку "
    "подтвердить перед подключением.",
    "Программы, где трафик из рассылок запрещён или требует согласия (Eventbrite), — только с письменным разрешением.",
]


def main() -> None:
    today = date.today().isoformat()
    aff = json.loads((ROOT / "data" / "affiliates_6d.json").read_text())
    tv = json.loads((ROOT / "data" / "ticket_vendors_6d.json").read_text())
    mon = json.loads((ROOT / "data" / "monetization_6d.json").read_text())
    shutil.copy(SRC, DST)
    wb = openpyxl.load_workbook(DST)
    ws = wb["Источники"]
    hdr = [c.value for c in ws[1]]
    if "Партнёрская программа" not in hdr:
        ws.cell(1, len(hdr) + 1, "Партнёрская программа")
        ws.cell(1, len(hdr) + 1).fill, ws.cell(1, len(hdr) + 1).font = HEADER_FILL, HEADER_FONT
        hdr.append("Партнёрская программа")
    col = hdr.index("Партнёрская программа") + 1
    con = connect()
    domains.REGISTRY = SRC   # хосты — по реестру v0.7 (с источниками этапа 6c)
    reg = domains.registry_hosts(con)
    events = {int(k): v for k, v in tv["events"].items()}
    by_src: dict[str, Counter] = {}
    for eid, v in events.items():
        for (sid,) in con.execute("SELECT DISTINCT source_id FROM event_sources WHERE event_id=?", (eid,)):
            by_src.setdefault(sid, Counter())[v["vendor"]] += 1
    src_hosts: dict[str, set] = {}
    for h, sids in reg.items():
        for sid in sids:
            src_hosts.setdefault(sid, set()).add(h)
    n_with = 0
    for r in range(2, ws.max_row + 1):
        sid = ws.cell(r, 1).value
        if not sid:
            continue
        own = next((a for a in aff["vendors"] for h in src_hosts.get(sid, ())
                    if any(h == x or h.endswith("." + x) for x in a["hosts"])), None)
        if own:
            text = f"{own['vendor']}: {own['programme']}" + (f" ({own['network']})" if own["network"] != "—" else "")
            n_with += has_programme(own)
        elif sid in by_src:
            parts = []
            for vend, n in by_src[sid].most_common(3):
                a = programme_of(vend, aff["vendors"])
                parts.append(f"продавец не определён — {n} соб." if vend == "не определён" else
                             f"{vend} — {n} соб.: " + ("программа есть" if has_programme(a) else "программы нет"
                                                       if a else "не исследовано"))
            text = "события ведут к: " + "; ".join(parts)
        else:
            text = ""
        ws.cell(r, col, text)
    ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 70
    for name in ("Монетизация",):
        if name in wb.sheetnames:
            del wb[name]
    m = wb.create_sheet("Монетизация")
    m.append(["Продавец", "Программа", "Сеть", "Комиссия", "Cookie", "Email разрешён", "Пометка", "Порог выплаты",
              "Источник условий", "Дата проверки"])
    for c in m[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for a in aff["vendors"]:
        m.append([a["vendor"], a["programme"], a["network"], a["commission"], a["cookie"], a["email"], a["disclosure"],
                  a["payout"], a["source"], aff["checked"]])
    m.append([])
    m.append(["Принципы подключения (зафиксированы на этапе 6d)"])
    m.cell(m.max_row, 1).font = Font(bold=True)
    for p in PRINCIPLES:
        m.append([p])
    m.append([])
    m.append(["Итоги (data/monetization_6d.json)"])
    m.cell(m.max_row, 1).font = Font(bold=True)
    w = mon["window"]
    m.append([f"События окна: {w['events']['n']}; продавец определён — {w['events']['vendor_known']:.0%}; ведут к продавцу "
              f"с программой — {w['events']['with_programme']:.0%} (с учётом важности — "
              f"{w['weighted_by_importance']['with_programme']:.0%})"])
    for ver, x in mon["items"].items():
        m.append([f"Пункты {ver}: {x['all']['n']}; к продавцу с программой — {x['all']['with_programme']:.0%} "
                  f"(с учётом важности — {x['weighted_by_importance']['with_programme']:.0%}); "
                  f"где email разрешён — {x['all']['with_programme_email_allowed']:.0%}"])
    for name, r in mon["estimate"]["result"].items():
        m.append([f"Доход на 1000 подписчиков в месяц, {name.replace('items_', 'по пунктам выпуска, ').replace('window_', 'по событиям окна, ')}: £{r['gbp_per_month_per_1000']} "
                  f"({r['clicks_per_month']} переходов, {r['orders_per_month']} заказов)"])
    for letter, width in zip("ABCDEFGHIJ", (34, 40, 16, 50, 16, 40, 40, 20, 60, 12)):
        m.column_dimensions[letter].width = width
    for row in m.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    # «Не разобрано» — свежая выгрузка
    if "Не разобрано" in wb.sheetnames:
        del wb["Не разобрано"]
    un = wb.create_sheet("Не разобрано")
    un.append(["ID", "Источник", "URL", "Тип проблемы", "Подробности", "С какой даты", "Что теряем", "Что делать",
               "Последняя проверка", "Статус"])
    for c in un[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    unparsed.init(con)
    for r in con.execute("SELECT * FROM unparsed_sources ORDER BY status, key"):
        un.append([r["key"], r["name"], r["url"], unparsed.PROBLEM_RU.get(r["problem"], r["problem"]), r["detail"],
                   r["since"], r["losing"], r["action"], (r["last_checked"] or "")[:10],
                   "открыто" if r["status"] == "open" else "решено"])
    for letter, width in zip("ABCDEFGHIJ", (22, 40, 50, 28, 50, 12, 45, 35, 12, 10)):
        un.column_dimensions[letter].width = width
    help_ws = wb["Как пользоваться"]
    help_ws.append([f"v0.8 ({today}): этап 6d — колонка «Партнёрская программа» и лист «Монетизация» (только оценка, "
                    "ни в какие программы не регистрировались); S167 — календарь релизов UK для рубрики «В кино»."])
    help_ws.cell(help_ws.max_row, 1).font = Font(italic=True)
    # S167 — календарь релизов
    ids = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
    if "S167" not in ids:
        ws.append(["S167", "Кино", "Календарь релизов UK (thepeoplesmovies.com + сверка mediamole.co.uk)", "Календарь",
                   "https://thepeoplesmovies.com/uk-release-dates-2026-2027", "", "Великобритания", "", "Еженедельно", 2])
        ws.cell(ws.max_row, hdr.index("Заметки") + 1 if "Заметки" in hdr else 11,
                "Этап 6d: новые фильмы недели для строки «В прокате с пятницы»; FDA launchingfilms.com — бот-защита, "
                "findanyfilm.com — заглушка, TMDB — нужен ключ (не регистрировались).")
    wb.save(DST)
    print(DST.relative_to(ROOT), "sources with own programme:", n_with)


if __name__ == "__main__":
    main()
