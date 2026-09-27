"""Этап 1: собирает data/cambridge_event_sources_v0.4.xlsx из v0.3 и результатов проверки.

Входы:
  data/cambridge_event_sources_v0.3.xlsx — исходный реестр (не изменяется);
  data/probe_results.json               — вывод scripts/probe_sources.py;
  data/p1_overrides.json                — ручные решения после разбора результатов (необязательно).

Выходы:
  data/cambridge_event_sources_v0.4.xlsx — реестр с обновлёнными строками приоритета 1
                                           и новым листом «Проверка P1»;
  docs/stage1_report.md                  — таблица «что нашёл / что не получилось».
"""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "cambridge_event_sources_v0.3.xlsx"
DST = ROOT / "data" / "cambridge_event_sources_v0.4.xlsx"
PROBE = ROOT / "data" / "probe_results.json"
OVERRIDES = ROOT / "data" / "p1_overrides.json"
REPORT = ROOT / "docs" / "stage1_report.md"

YELLOW = PatternFill("solid", fgColor="FFFFF2CC")
RED = PatternFill("solid", fgColor="FFF4CCCC")
GREEN = PatternFill("solid", fgColor="FFD9EAD3")
NO_FILL = PatternFill(fill_type=None)
HEADER_FILL = PatternFill("solid", fgColor="FF1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")

# Порядок предпочтения способов сбора (чем меньше, тем лучше).
RANK = {"iCal": 0, "RSS": 1, "API": 2, "JSON-LD": 3, "HTML": 4}


def load_json(path: Path) -> dict:
    return json.loads(path.read_text()) if path.exists() else {}


def summarize(probe: dict) -> dict:
    """Сводит проверку всех URL источника в одно решение."""
    urls = probe.get("urls", [])
    live = [u for u in urls if u.get("status") == 200]
    blocked = [u for u in urls if u.get("robots") == "запрещено"]
    failed = [u for u in urls if u.get("robots") != "запрещено" and u.get("status") != 200]

    methods: list[tuple[str, str]] = []  # (метод, endpoint)
    for u in live:
        url = u["final_url"]
        kind = u.get("kind")
        if kind == "ical" and u.get("ical_events", 0) >= 0:
            methods.append(("iCal", url))
        elif kind == "feed":
            methods.append(("RSS", url))
        elif kind == "json":
            methods.append(("API", url))
        elif kind == "html":
            for f in u.get("alternate_links", []):
                if "calendar" in f["type"]:
                    methods.append(("iCal", f["url"]))
                elif ("rss" in f["type"] or "atom" in f["type"]) and "/comments/feed" not in f["url"]:
                    methods.append(("RSS", f["url"]))  # JSON-альтернативы (oEmbed, wp/v2/pages) — не фиды событий
            for link in u.get("ical_links", []):
                methods.append(("iCal", link))
            if u.get("jsonld_events"):
                methods.append(("JSON-LD", url))
            methods.append(("HTML", url))
    methods.sort(key=lambda m: RANK[m[0]])

    found, problems = [], []
    for u in live:
        bits = [f"{u['final_url']} — 200"]
        if u.get("kind") == "feed":
            bits.append(f"фид, записей: {u.get('feed_entries')}")
        if u.get("kind") == "ical":
            bits.append(f"iCal, событий: {u.get('ical_events')}")
        if u.get("kind") == "json":
            bits.append("JSON")
        if u.get("jsonld_events"):
            bits.append(f"JSON-LD Event: {u['jsonld_events']}")
        if u.get("alternate_links"):
            bits.append("alternate-фиды: " + ", ".join(f["url"] for f in u["alternate_links"][:3]))
        if u.get("tribe_events"):
            bits.append("WordPress + The Events Calendar")
        elif u.get("wordpress"):
            bits.append("WordPress")
        if u.get("mentions"):
            bits.append("билетная система: " + ", ".join(u["mentions"]))
        found.append("; ".join(bits))
    for u in blocked:
        problems.append(f"{u['url']} — запрещено robots.txt")
    for u in failed:
        problems.append(f"{u['url']} — {u.get('status') or u.get('error') or u.get('robots') or 'нет ответа'}")

    return {
        "live": bool(live),
        "best": methods[0] if methods else None,
        "found": found,
        "problems": problems,
        "robots": "запрещено" if blocked and not live else ("разрешено" if live else "—"),
    }


def main() -> None:
    probes = load_json(PROBE)
    overrides = {k: v for k, v in load_json(OVERRIDES).items() if not k.startswith("_")}
    today = date.today().isoformat()

    wb = openpyxl.load_workbook(SRC)
    ws = wb["Источники"]

    new_cols = ["Endpoint для сбора", "robots.txt", "Проверено"]
    first_new = ws.max_column + 1
    for i, name in enumerate(new_cols):
        c = ws.cell(row=1, column=first_new + i, value=name)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[openpyxl.utils.get_column_letter(first_new)].width = 50
    ws.column_dimensions[openpyxl.utils.get_column_letter(first_new + 1)].width = 12
    ws.column_dimensions[openpyxl.utils.get_column_letter(first_new + 2)].width = 12
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(first_new + 2)}{ws.max_row}"

    for dv in ws.data_validations.dataValidation:
        if "K2" in str(dv.sqref):
            dv.formula1 = '"да,проверить,нет"'

    rows_report = []
    for row in ws.iter_rows(min_row=2):
        sid, name, url, method, prio, verified, notes = (row[0].value, row[2].value, row[4], row[7], row[9].value, row[10], row[11])
        if prio != 1:
            continue
        ov = overrides.get(sid, {})
        s = summarize(probes.get(sid, {})) if sid in probes else None

        new_url = ov.get("url")
        new_method = ov.get("method")
        endpoint = ov.get("endpoint")
        status = ov.get("verified")
        if s:
            if not new_url and s["live"]:
                new_url = next(u["final_url"] for u in probes[sid]["urls"] if u.get("status") == 200)
            if s["best"]:
                new_method = new_method or s["best"][0]
                endpoint = endpoint or s["best"][1]
            status = status or ("да" if s["live"] else "нет")

        if new_url:
            url.value = new_url
        if new_method:
            method.value = new_method
        if status:
            verified.value = status
        note = ov.get("note")
        if note:
            notes.value = f"{notes.value}. v0.4: {note}" if notes.value else f"v0.4: {note}"
        ws.cell(row=row[0].row, column=first_new, value=endpoint)
        ws.cell(row=row[0].row, column=first_new + 1, value=ov.get("robots") or (s["robots"] if s else None))
        ws.cell(row=row[0].row, column=first_new + 2, value=today if (s or ov) else None)

        fill = {"да": NO_FILL, "нет": RED}.get(verified.value, YELLOW)
        url.fill = fill if url.value else YELLOW
        verified.fill = fill

        rows_report.append({
            "id": sid, "name": name, "status": verified.value, "method": method.value,
            "endpoint": endpoint or "", "url": url.value or "",
            "found": ov.get("found") or (s["found"] if s else []),
            "problems": ov.get("problems") or (s["problems"] if s else ["не проверялся"]),
        })

    # Лист «Проверка P1»
    if "Проверка P1" in wb.sheetnames:
        del wb["Проверка P1"]
    rep = wb.create_sheet("Проверка P1", index=wb.sheetnames.index("Источники") + 1)
    header = ["ID", "Источник", "URL проверен", "Способ сбора", "Endpoint", "Что нашёл", "Что не получилось"]
    rep.append(header)
    for c in rep[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for r in rows_report:
        rep.append([r["id"], r["name"], r["status"], r["method"], r["endpoint"],
                    "\n".join(r["found"]), "\n".join(r["problems"])])
        fill = {"да": GREEN, "нет": RED}.get(r["status"], YELLOW)
        rep.cell(row=rep.max_row, column=3).fill = fill
    for col, width in zip("ABCDEFG", (7, 36, 12, 18, 50, 70, 50)):
        rep.column_dimensions[col].width = width
    for row in rep.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    rep.freeze_panes = "C2"
    rep.auto_filter.ref = f"A1:G{rep.max_row}"

    ws_help = wb["Как пользоваться"]
    ws_help.cell(row=1, column=1, value=str(ws_help["A1"].value).replace("v0.3, сентябрь 2026", "v0.4, сентябрь 2026"))
    ws_help.append([f"v0.4 ({today}): проверены источники приоритета 1 — см. лист «Проверка P1»; новые столбцы "
                    "«Endpoint для сбора», «robots.txt», «Проверено». «URL проверен = нет» — URL не отвечает (красный)."])

    wb.save(DST)

    lines = ["# Этап 1 — проверка источников приоритета 1", "",
             f"Дата проверки: {today}. Реестр: `data/cambridge_event_sources_v0.4.xlsx`.", "",
             "| ID | Источник | Проверен | Способ | Что нашёл | Что не получилось |",
             "|---|---|---|---|---|---|"]
    esc = lambda t: str(t).replace("|", "\\|").replace("\n", "<br>")  # noqa: E731
    for r in rows_report:
        lines.append(f"| {r['id']} | {esc(r['name'])} | {r['status']} | {esc(r['method'])} | "
                     f"{esc('<br>'.join(r['found'])) or '—'} | {esc('<br>'.join(r['problems'])) or '—'} |")
    REPORT.write_text("\n".join(lines) + "\n")
    print(f"Сохранено: {DST}\nОтчёт: {REPORT}")


if __name__ == "__main__":
    main()
