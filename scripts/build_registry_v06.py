"""Этап 6: data/cambridge_event_sources_v0.6.xlsx из v0.5 — HTML-источники P1, семейные, поля/церкви/ярмарки,
отложенные P2 и P3.

Входы: v0.5 (не изменяется), data/p6_decisions.json. Строки получают решение этапа 6 в «Заметках» и цвет ID; новые
источники (S130–S147) добавляются, «Семейные места и аттракционы» — новая категория с приоритетом 2; лист «Этап 6» —
таблица «источник → решение → новых уникальных событий» (числа — из docs/stage6_unique.json, его пишет отчёт).
"""

from __future__ import annotations

import json
import shutil
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from build_registry_v04 import GREEN, HEADER_FILL, HEADER_FONT, RED, YELLOW  # noqa: E402

SRC = ROOT / "data" / "cambridge_event_sources_v0.5.xlsx"
DST = ROOT / "data" / "cambridge_event_sources_v0.6.xlsx"
FILL = {"подключён": GREEN, "закрыт": RED, "отложен": YELLOW, "вручную": YELLOW}
BLOCK = {"P1": "HTML-источники P1", "family": "Семейные источники", "places": "Семейные места и аттракционы",
         "open": "Поля, церкви, ярмарки", "P2": "Отложенные P2", "P3": "Источники P3"}
# новые источники этапа 6: название, тип, URL, категория, зона
NEW = {
    "S130": ("Museum of Cambridge", "Музей", "https://www.museumofcambridge.org.uk/events/", "Семейные источники", "Центр"),
    "S131": ("Cambridge Past, Present & Future (Wandlebury, Leper Chapel, Hinxton Watermill…)", "Организатор",
             "https://cambridgeppf.org/events-list/", "Семейные источники", "До 30 мин"),
    "S132": ("Milton Country Park", "Место", "https://www.miltoncountrypark.org/events", "Семейные места и аттракционы", "До 30 мин"),
    "S133": ("Nene Park / Ferry Meadows", "Место", "https://www.nenepark.org.uk/events/", "Семейные места и аттракционы", "Кембриджшир"),
    "S134": ("Centre for Computing History", "Музей", "https://www.computinghistory.org.uk/pages/30677/What-s-On/", "Семейные источники", "Центр"),
    "S135": ("Shepreth Wildlife Park", "Место", "https://www.sheprethwildlifepark.co.uk", "Семейные места и аттракционы", "До 30 мин"),
    "S136": ("Linton Zoo", "Место", "https://lintonzoo.com", "Семейные места и аттракционы", "До 30 мин"),
    "S137": ("Hamerton Zoo Park", "Место", "https://www.hamertonzoopark.com", "Семейные места и аттракционы", "До часа"),
    "S138": ("Wimpole Home Farm (National Trust)", "Место", "https://www.nationaltrust.org.uk/visit/cambridgeshire/wimpole",
             "Семейные места и аттракционы", "До 30 мин"),
    "S139": ("Audley End House (English Heritage)", "Место",
             "https://www.english-heritage.org.uk/visit/places/audley-end-house-and-gardens/events/",
             "Семейные места и аттракционы", "До 30 мин"),
    "S140": ("Woburn Safari Park", "Место", "https://www.woburnsafari.co.uk", "Семейные места и аттракционы", "40–60 км"),
    "S141": ("A Church Near You (Church of England)", "Каталог", "https://www.achurchnearyou.com", "Церкви", "Графство"),
    "S142": ("ChurchSuite", "Каталог", "https://churchsuite.com", "Церкви", "Графство"),
    "S143": ("King's College Chapel", "Площадка", "https://www.kings.cam.ac.uk/chapel/music-chapel/concerts-kings", "Церкви", "Центр"),
    "S144": ("St John's College Chapel", "Площадка", "https://www.joh.cam.ac.uk/chapel-and-choir/chapel-services-and-events",
             "Церкви", "Центр"),
    "S145": ("Trinity College Chapel / Great St Mary's", "Площадка", "https://www.greatstmarys.org/whats-on", "Церкви", "Центр"),
    "S146": ("Round Church / St Botolph's", "Площадка", "https://roundchurchcambridge.org", "Церкви", "Центр"),
    "S147": ("Stourbridge Fair, Big Weekend, включение огней", "Ежегодное", "https://cambridgeppf.org/leper-chapel/",
             "Ярмарки", "Центр"),
}


def main() -> None:
    today = date.today().isoformat()
    dec = {k: v for k, v in json.loads((ROOT / "data" / "p6_decisions.json").read_text()).items() if not k.startswith("_")}
    uniq_path = ROOT / "docs" / "stage6_unique.json"
    uniq = json.loads(uniq_path.read_text()) if uniq_path.exists() else {}
    shutil.copy(SRC, DST)
    wb = openpyxl.load_workbook(DST)
    ws = wb["Источники"]
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr)}
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    report = []
    for sid, (state, decision, note, block) in dec.items():
        if sid not in rows:
            nm = NEW.get(sid, (sid, "", None, BLOCK[block], ""))
            prio = 2 if block in ("places", "family", "open") else 3
            ws.append([sid, nm[3], nm[0], nm[1], nm[2], "", nm[4], "", "Ежедневно", prio])
            rows[sid] = ws.max_row
        r = rows[sid]
        old = ws.cell(r, col["Заметки"]).value or ""
        ws.cell(r, col["Заметки"], (old + " " if old else "") + f"Этап 6: {decision} — {note}")
        ws.cell(r, col["Проверено"], today)
        ws.cell(r, col["URL проверен"], "нет" if state == "закрыт" else "да" if state == "живой" else state)
        ws.cell(r, col["ID"]).fill = FILL.get(decision, PatternFill(fill_type=None))
        report.append([sid, ws.cell(r, col["Источник"]).value, BLOCK[block], state, decision,
                       uniq.get(sid, {}).get("unique", "—"), note])
    if "Этап 6" in wb.sheetnames:
        del wb["Этап 6"]
    rep = wb.create_sheet("Этап 6")
    rep.append(["ID", "Источник", "Блок", "Состояние", "Решение", "Новых уникальных (будущие, в зоне)", "Примечание"])
    for c in rep[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    order = list(BLOCK.values())
    for row in sorted(report, key=lambda x: (order.index(x[2]), x[0])):
        rep.append(row)
        rep.cell(rep.max_row, 5).fill = FILL.get(row[4], PatternFill(fill_type=None))
    for letter, width in zip("ABCDEFG", (8, 45, 28, 12, 14, 18, 100)):
        rep.column_dimensions[letter].width = width
    help_ws = wb["Как пользоваться"]
    help_ws.append([f"v0.6 ({today}): этап 6 — HTML-источники P1, семейные источники и новая категория «Семейные места "
                    "и аттракционы» (P2), поля/церкви/ярмарки, отложенные P2, P3; см. лист «Этап 6»."])
    help_ws.cell(help_ws.max_row, 1).font = Font(italic=True)
    wb.save(DST)
    print(DST.relative_to(ROOT), len(report))


if __name__ == "__main__":
    main()
