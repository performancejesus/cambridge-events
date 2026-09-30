"""Этап 7b: data/cambridge_event_sources_v0.9.xlsx из v0.8 (v0.8 не изменяется).

- Лист «Источники»: Cambridge Union (S049) и National Garden Scheme (S067) подключены; Light Cinema (S157) — открытый
  JSON мини-гида; Trinity College (S166) — публичные лекции; новые S168 (Camdram — все площадки) и S169–S175
  (страницы событий колледжей).
- Лист «Колледжи»: все 31 колледж — страница, что там, фид, robots.txt, события в базе до и после, решение
  (data/college_coverage_7b.json).
- Лист «Не разобрано» — свежая выгрузка unparsed_sources (с закрытыми сайтами колледжей и хоров).
"""

from __future__ import annotations

import json
import shutil
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from build_registry_v04 import HEADER_FILL, HEADER_FONT  # noqa: E402
from pipeline import unparsed  # noqa: E402
from pipeline.db import connect  # noqa: E402

SRC = ROOT / "data" / "cambridge_event_sources_v0.8.xlsx"
DST = ROOT / "data" / "cambridge_event_sources_v0.9.xlsx"
T = "2026-09-30"
UPDATES = {   # ID → (URL, Endpoint, способ сбора, заметка этапа 7b)
    "S049": ("https://cus.org/event", "https://cus.org/wp-json/wp/v2/event_jet", "JSON (WordPress REST) + страница события",
             "Этап 7b: подключён — события EventJet — записи WordPress event_jet (открытый REST, без токенов); уровень "
             "доступа и площадка — со страницы события, дата и цена — Haiku. Доступ: Open To The Public / Ticketed — для "
             "всех; Members Only, Members & Cambridge Students — «только для членов», членство Open для всех £370 в год; "
             "Life/Annual/Access, Students, Staff — не берём. Termcard — Issuu (картинки), не разбираем. На 30.09 "
             "опубликованных событий 0 (Michaelmas не выложен). cus.org отвечает 429 — пауза 10 с. The Orator's Cellar на "
             "Skiddle — «No events to display»; площадки The Orator / The Footlight Cellars склеиваются синонимами."),
    "S067": ("https://findagarden.ngs.org.uk", "https://api.findagarden.ngs.org.uk/api/gardens", "JSON (API приложения)",
             "Этап 7b: подключён — открытый API приложения findagarden (robots.txt разрешает): сады Кембриджшира и в 75 км, "
             "открытия с датой и ценой; подряд идущие дни — одним событием (Robinson College — ежедневно до 22 декабря)."),
    "S157": ("https://cambridge.thelight.co.uk", "https://cambridge.thelight.co.uk/resource/services/miniguide/data.ashx",
             "JSON (мини-гид)", "Этап 7b: подключён — расписание на 2 недели открытым JSON без токена (robots.txt "
             "разрешает): event cinema и спецпоказы — события; фильмы — «где идёт» в рубрике «В кино»."),
    "S166": ("https://www.trin.cam.ac.uk/events/", None, "HTML + Haiku",
             "Этап 7b: подключён — публичные лекции (Birkbeck Lectures, /about/public-lectures/); встречи выпускников "
             "отсекаются (решение этапа 6c «не нужен» пересмотрено)."),
}
NEW = [  # ID, категория, источник, тип, URL, покрывает, заметка
    ("S168", "Театр", "Camdram — все площадки", "Агрегатор", "https://www.camdram.net/diary.json",
     "Corpus Playroom, театры колледжей (Fitzpatrick Hall, Robinson Brickhouse, Pembroke New Cellars, Howard Theatre)",
     "Этап 7b: общий дневник Camdram (JSON), площадки ADC — S042."),
    ("S169", "Колледжи", "Churchill College + Churchill Archives Centre", "Колледж", "https://www.chu.cam.ac.uk/events/list/",
     "публичные лекции и события", "Этап 7b: Haiku по странице событий; выпускники и студенты — не берём."),
    ("S170", "Колледжи", "Clare Hall", "Колледж", "https://www.clarehall.cam.ac.uk/events/",
     "лекции, концерты, выставки", "Этап 7b"),
    ("S171", "Колледжи", "Lucy Cavendish College", "Колледж", "https://www.lucy.cam.ac.uk/events", "лекции, панели", "Этап 7b"),
    ("S172", "Колледжи", "Robinson College", "Колледж", "https://www.robinson.cam.ac.uk/events", "концерты",
     "Этап 7b: внутренние занятия для студентов отсекаются"),
    ("S173", "Колледжи", "St Edmund's College — Von Hügel Institute", "Колледж",
     "https://www.st-edmunds.cam.ac.uk/the-vhi/vhi-events/", "публичные лекции VHI", "Этап 7b"),
    ("S174", "Колледжи", "Girton College", "Колледж", "https://www.girton.cam.ac.uk/upcoming-events", "концерты, лекции",
     "Этап 7b"),
    ("S175", "Колледжи", "Hughes Hall", "Колледж", "https://www.hughes.cam.ac.uk/about/events/",
     "публичные лекции, ярмарки", "Этап 7b: большинство событий — для студентов, отсекаются"),
]


def main() -> None:
    shutil.copy(SRC, DST)
    wb = openpyxl.load_workbook(DST)
    ws = wb["Источники"]
    hdr = [c.value for c in ws[1]]
    ci = {h: hdr.index(h) + 1 for h in hdr if h}
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
    for sid, (url, endpoint, how, note) in UPDATES.items():
        r = rows[sid]
        ws.cell(r, ci["URL"], url)
        if endpoint:
            ws.cell(r, ci["Endpoint для сбора"], endpoint)
        ws.cell(r, ci["Способ сбора"], how)
        ws.cell(r, ci["Частота"], "Ежедневно")
        ws.cell(r, ci["Заметки"], ((ws.cell(r, ci["Заметки"]).value or "") + " " + note).strip())
        ws.cell(r, ci["robots.txt"], "разрешено")
        ws.cell(r, ci["Проверено"], T)
    for sid, cat, name, typ, url, covers, note in NEW:
        if sid in rows:
            continue
        ws.append([None] * len(hdr))
        r = ws.max_row
        for k, v in {"ID": sid, "Категория": cat, "Источник": name, "Тип источника": typ, "URL": url,
                     "Что покрывает": covers, "Зона": "Центр", "Способ сбора": "JSON" if url.endswith(".json") else "HTML + Haiku",
                     "Частота": "Ежедневно", "Приоритет": 2, "URL проверен": "да", "Заметки": note,
                     "robots.txt": "разрешено", "Проверено": T}.items():
            ws.cell(r, ci[k], v)
    # «Колледжи»
    cov = json.loads((ROOT / "data" / "college_coverage_7b.json").read_text())
    if "Колледжи" in wb.sheetnames:
        del wb["Колледжи"]
    k = wb.create_sheet("Колледжи")
    k.append(["Колледж", "Страница публичных событий", "Что там", "Фид", "robots.txt", "До: событий за 60 дн. назад / вперёд",
              "После: за 60 дн. назад / вперёд", "Источники", "Решение"])
    for c in k[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for x in cov["colleges"]:
        k.append([x["college"], x["page"], x["what"], x["feed"], x["robots"],
                  f"{x['before']['past_60']} / {x['before']['next_60']}", f"{x['after']['past_60']} / {x['after']['next_60']}",
                  ", ".join(sorted(x["after"]["sources"])) or "—", x["decision"]])
    s = cov["summary"]
    k.append([])
    k.append([f"Покрыто (есть будущие события в базе): до — {s['covered_before']} из 31, после — {s['covered_after']} из 31; "
              f"свой коллектор — {s['own_collector']}; закрыты бот-защитой или обрывом — {s['blocked']}; публичной афиши "
              f"нет — {s['no_public_programme']}."])
    for letter, width in zip("ABCDEFGHI", (18, 42, 50, 24, 12, 16, 16, 30, 60)):
        k.column_dimensions[letter].width = width
    for row in k.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    # «Не разобрано»
    if "Не разобрано" in wb.sheetnames:
        del wb["Не разобрано"]
    un = wb.create_sheet("Не разобрано")
    un.append(["ID", "Источник", "URL", "Тип проблемы", "Подробности", "С какой даты", "Что теряем", "Что делать",
               "Последняя проверка", "Статус"])
    for c in un[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    con = connect()
    unparsed.init(con)
    for r in con.execute("SELECT * FROM unparsed_sources ORDER BY status, key"):
        un.append([r["key"], r["name"], r["url"], unparsed.PROBLEM_RU.get(r["problem"], r["problem"]), r["detail"],
                   r["since"], r["losing"], r["action"], (r["last_checked"] or "")[:10],
                   "открыто" if r["status"] == "open" else "решено"])
    for letter, width in zip("ABCDEFGHIJ", (22, 40, 50, 28, 50, 12, 45, 35, 12, 10)):
        un.column_dimensions[letter].width = width
    help_ws = wb["Как пользоваться"]
    help_ws.append([f"v0.9 ({T}): этап 7b — Cambridge Union, NGS, Light Cinema подключены; Camdram — все площадки; "
                    "страницы событий 7 колледжей (S169–S175) и Trinity (S166); лист «Колледжи» — аудит всех 31."])
    help_ws.cell(help_ws.max_row, 1).font = Font(italic=True)
    wb.save(DST)
    print(DST.relative_to(ROOT))


if __name__ == "__main__":
    main()
