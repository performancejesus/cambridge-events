"""Этап 5: data/cambridge_event_sources_v0.5.xlsx из v0.4 — проверка P2 и новые источники зоны.

Входы: v0.4 (не изменяется), data/probe_results_p2.json, data/p2_candidates.json (названия новых источников),
data/p2_decisions.json (решения). Строки P2 получают endpoint, robots.txt, дату проверки и решение в «Заметках»;
новые источники S098–S129 добавляются с приоритетом 2; лист «Проверка P2» — таблица проверки.
"""

from __future__ import annotations

import json
import shutil
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from build_registry_v04 import GREEN, HEADER_FILL, HEADER_FONT, RED, YELLOW, summarize  # noqa: E402

SRC = ROOT / "data" / "cambridge_event_sources_v0.4.xlsx"
DST = ROOT / "data" / "cambridge_event_sources_v0.5.xlsx"
FILL = {"подключён": GREEN, "закрыт": RED, "отложен": YELLOW, "нужно решение": YELLOW}


def main() -> None:
    today = date.today().isoformat()
    probe = json.loads((ROOT / "data" / "probe_results_p2.json").read_text())
    names = json.loads((ROOT / "data" / "p2_candidates.json").read_text())["_names"]
    dec = {k: v for k, v in json.loads((ROOT / "data" / "p2_decisions.json").read_text()).items() if not k.startswith("_")}
    shutil.copy(SRC, DST)
    wb = openpyxl.load_workbook(DST)
    ws = wb["Источники"]
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr)}
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    report = []
    for sid, (state, decision, note) in dec.items():
        key = {"S128": "S006x", "S129": "S007x"}.get(sid, sid)   # страницы городов проверялись как S006x / S007x
        s = summarize(probe[key]) if key in probe else None
        method, endpoint = (s.get("best") or (None, None)) if s else (None, None)
        robots = s["robots"] if s else None
        if sid not in rows:  # новый источник зоны
            nm = names.get(sid, [sid, "", ""])
            ws.append([sid, "Расширение географии (этап 5)", nm[0], nm[1], (probe.get(key, {}).get("urls") or [{}])[0].get("url"),
                       nm[2], "Час / Кембриджшир", method, "Ежедневно", 2])
            rows[sid] = ws.max_row
        r = rows[sid]
        ws.cell(r, col["URL проверен"], "да" if state in ("живой", "частично") else "нет" if state == "закрыт" else state)
        ws.cell(r, col["Endpoint для сбора"], endpoint)
        ws.cell(r, col["robots.txt"], robots)
        ws.cell(r, col["Проверено"], today)
        ws.cell(r, col["Заметки"], f"Этап 5: {decision} — {note}")
        ws.cell(r, col["ID"]).fill = FILL.get(decision, PatternFill(fill_type=None))
        report.append([sid, ws.cell(r, col["Источник"]).value, state, decision, method,
                       endpoint, note])
    if "Проверка P2" in wb.sheetnames:
        del wb["Проверка P2"]
    rep = wb.create_sheet("Проверка P2")
    rep.append(["ID", "Источник", "Состояние", "Решение", "Способ", "Endpoint", "Примечание"])
    for c in rep[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for row in sorted(report):
        rep.append(row)
        rep.cell(rep.max_row, 4).fill = FILL.get(row[3], PatternFill(fill_type=None))
    for letter, width in zip("ABCDEFG", (8, 40, 12, 16, 14, 60, 90)):
        rep.column_dimensions[letter].width = width
    help_ws = wb["Как пользоваться"]
    help_ws.append([f"v0.5 ({today}): этап 5 — проверены источники приоритета 2 и новые источники зоны (S098–S129), "
                    "см. лист «Проверка P2» и столбец «Заметки»; зелёный — подключён, жёлтый — отложен / нужно решение, "
                    "красный — закрыт."])
    help_ws.cell(help_ws.max_row, 1).font = Font(italic=True)
    wb.save(DST)
    print(DST.relative_to(ROOT), len(report))


if __name__ == "__main__":
    main()
